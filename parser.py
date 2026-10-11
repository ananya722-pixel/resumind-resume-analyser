
"""
parser.py

Extracts plain text from uploaded PDF or DOCX files.

This is the only module responsible for reading resume files from disk.
"""

import os
import re

import docx
import pdfplumber


def parse_resume(file_path: str) -> str:
    """
    Detect file type by extension and route to the appropriate parser.

    Returns:
        Clean plain-text content extracted from the resume.

    Raises:
        ValueError: If the file type is unsupported or no readable text
                    can be extracted.
    """
    if not os.path.isfile(file_path):
        raise FileNotFoundError(f"Resume file not found: {file_path}")

    ext = os.path.splitext(file_path)[1].lower()

    if ext == ".pdf":
        text = _parse_pdf(file_path)
    elif ext == ".docx":
        text = _parse_docx(file_path)
    else:
        raise ValueError(
            f"Unsupported file extension: {ext}. "
            "Please upload a PDF or DOCX file."
        )

    cleaned = _clean_text(text)

    if len(cleaned.strip()) < 30:
        raise ValueError(
            "No readable text could be extracted from this resume. "
            "The PDF may be scanned or image-only. "
            "Try uploading a text-based PDF or a DOCX file."
        )

    return cleaned


# ── Private helpers ─────────────────────────────────────────────


def _parse_pdf(file_path: str) -> str:
    """
    Extract text from every page of a PDF using pdfplumber.
    """
    text_pages = []

    with pdfplumber.open(file_path) as pdf:
        if not pdf.pages:
            raise ValueError("The PDF contains no pages.")

        for page in pdf.pages:
            page_text = page.extract_text()

            if page_text:
                text_pages.append(page_text)

    return "\n".join(text_pages)


def _parse_docx(file_path: str) -> str:
    """
    Extract text from DOCX paragraphs and tables.
    """
    document = docx.Document(file_path)
    sections = []

    # Extract paragraphs
    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            sections.append(paragraph.text)

    # Extract table content, which often contains resume details
    for table in document.tables:
        for row in table.rows:
            cells = [
                cell.text.strip()
                for cell in row.cells
                if cell.text.strip()
            ]

            if cells:
                sections.append(" | ".join(cells))

    return "\n".join(sections)


def _clean_text(text: str) -> str:
    """
    Remove null bytes and excessive blank lines while preserving
    line breaks and resume structure.
    """
    text = text.replace("\x00", "")

    # Collapse 3 or more consecutive newlines into 2
    text = re.sub(r"\n{3,}", "\n\n", text)

    # Remove trailing whitespace from each line
    lines = [line.rstrip() for line in text.splitlines()]

    return "\n".join(lines).strip()

